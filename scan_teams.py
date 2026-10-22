"""Quet danh sach doi tu cot Doi (openpyxl, khong can mo Excel GUI)."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent


def load_config() -> dict:
    import json

    with (ROOT / "config.json").open(encoding="utf-8") as f:
        return json.load(f)


def scan_teams_from_workbook(wb_path: Path) -> list[str]:
    cfg = load_config()
    header_row = int(cfg.get("header_row", 2))
    filter_col = cfg.get("filter_column", "Đội")
    skip = set(cfg.get("skip_teams", []))
    sheet_names = [s["name"] for s in cfg["sheets"]]

    found: set[str] = set()
    wb = load_workbook(wb_path, read_only=True, data_only=True)
    try:
        for sheet_name in sheet_names:
            if sheet_name not in wb.sheetnames:
                raise ValueError(f"Thieu sheet: {sheet_name}")
            ws = wb[sheet_name]
            header_cells = next(
                ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True)
            )
            doi_idx = None
            for j, h in enumerate(header_cells):
                if h is not None and str(h).strip() == filter_col:
                    doi_idx = j
                    break
            if doi_idx is None:
                raise ValueError(f"Khong tim thay cot '{filter_col}' o sheet {sheet_name}")

            for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
                if doi_idx >= len(row):
                    continue
                val = row[doi_idx]
                if val is None:
                    continue
                name = str(val).strip()
                if name and name not in skip:
                    found.add(name)
    finally:
        wb.close()

    return sorted(found)
